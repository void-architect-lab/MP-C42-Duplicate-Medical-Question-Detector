"""Generate the MP-C42 report from real outputs produced by train.py.

The document leaves student identity and public URLs editable, and appends the
actual Python source code. No performance metrics are fabricated.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs"
RESULTS = OUT / "results.json"
REPORT = OUT / "MP-C42_Report.docx"


def _fmt_pct(value: Any) -> str:
    try:
        return f"{float(value) * 100:.2f}%"
    except (TypeError, ValueError):
        return "N/A"


def _fmt_num(value: Any, digits: int = 4) -> str:
    try:
        return f"{float(value):.{digits}f}"
    except (TypeError, ValueError):
        return "N/A"


def _shade(cell, fill: str = "D9EAF7") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_text(cell, text: Any, bold: bool = False, size: float = 8.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _table(doc: Document, headers: list[str], rows: list[list[Any]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, head in enumerate(headers):
        _set_cell_text(table.rows[0].cells[idx], head, bold=True, size=8.5)
        _shade(table.rows[0].cells[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            _set_cell_text(cells[idx], value, size=8)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def _body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.04


def _bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.add_run(text)


def _add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = paragraph.add_run("Page ")
    r.font.name = "Times New Roman"
    r.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def _add_image(doc: Document, path: Path, caption: str, width: float = 5.7) -> None:
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        cp = doc.add_paragraph(caption)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(4)
        for run in cp.runs:
            run.italic = True
            run.font.size = Pt(8.5)
    else:
        _body(doc, f"[{caption}: image not found. Rerun training to generate it.]")


def _add_source_appendix(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("APPENDIX A — COMPLETE SOURCE CODE", level=1)
    _body(doc, "The following is the project source used to reproduce the pipeline. Generated weights and dataset contents are intentionally not included. Keep the public GitHub URL and notebook URL updated before submission.")
    source_files = [
        ROOT / "src" / "data_utils.py",
        ROOT / "src" / "gpt_model.py",
        ROOT / "train.py",
        ROOT / "app.py",
        ROOT / "offline_demo.py",
        ROOT / "build_report.py",
    ]
    for path in source_files:
        if not path.exists():
            continue
        doc.add_heading(f"A.3.{source_files.index(path)+1} {path.relative_to(ROOT).as_posix()}", level=2)
        code = path.read_text(encoding="utf-8").splitlines()
        for line in code:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_together = True
            run = p.add_run(line if line else " ")
            run.font.name = "Courier New"
            run.font.size = Pt(6.5)


def build_report(root: Path = ROOT, require_results: bool = True) -> Path:
    global ROOT, OUT, RESULTS, REPORT
    ROOT = Path(root).resolve()
    OUT = ROOT / "outputs"
    RESULTS = OUT / "results.json"
    REPORT = OUT / "MP-C42_Report.docx"
    if require_results and not RESULTS.exists():
        raise FileNotFoundError(f"Training results not found at {RESULTS}; run train.py first.")
    if not RESULTS.exists():
        raise FileNotFoundError("Run train.py first so the report uses measured model results.")
    r = json.loads(RESULTS.read_text(encoding="utf-8"))
    before = r["before"]
    after = r["after"]
    splits = r["split_sizes"]
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for style_name, size, color in [("Title", 21, "17365D"), ("Heading 1", 14, "17365D"), ("Heading 2", 11, "1F4E79"), ("Heading 3", 10, "1F4E79")]:
        st = styles[style_name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
    header = section.header.paragraphs[0]
    header.text = "Large Language Models (ISE23704C) | MP-C42"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for rr in header.runs:
        rr.font.size = Pt(8)
        rr.font.italic = True
    _add_page_field(section.footer.paragraphs[0])

    # Cover page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    run = p.add_run("GLOBAL ACADEMY OF TECHNOLOGY")
    run.bold = True; run.font.name = "Times New Roman"; run.font.size = Pt(17)
    p = doc.add_paragraph("Autonomous Institute, Affiliated to Visvesvaraya Technological University, Belagavi\nRajarajeshwari Nagar, Bengaluru — 560 098")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs: run.font.size = Pt(10); run.italic = True
    doc.add_paragraph()
    p = doc.add_paragraph("A Mini Project Report on")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("DUPLICATE MEDICAL QUESTION DETECTOR")
    run.bold = True; run.font.size = Pt(18); run.font.color.rgb = RGBColor(31,78,121)
    p = doc.add_paragraph("Project ID: MP-C42")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Submitted in partial fulfilment of the requirements of\nLARGE LANGUAGE MODELS (ISE23704C)\nVII Semester, Bachelor of Engineering")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Submitted by\n[STUDENT NAME]  |  [USN]  |  [SECTION 7A/7B]\n\nSDG 3 — Good Health and Well-being\nSDG 9 — Industry, Innovation and Infrastructure\n\nUnder the guidance of\n[COURSE FACULTY / GUIDE NAME]")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    p = doc.add_paragraph("DEPARTMENT OF INFORMATION SCIENCE AND ENGINEERING\nGLOBAL ACADEMY OF TECHNOLOGY, BENGALURU\nAcademic Year 2026–27")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(34)
    for run in p.runs: run.bold = True; run.font.size = Pt(10)

    doc.add_page_break()
    doc.add_heading("CERTIFICATE", level=1)
    _body(doc, "This is to certify that the mini project entitled “Duplicate Medical Question Detector” (Project ID MP-C42) is a bonafide work carried out by [STUDENT NAME], USN [USN], of VII Semester, Section [7A/7B], in partial fulfilment of the requirements of the course Large Language Models (ISE23704C) during the academic year 2026–27.")
    doc.add_paragraph("\n\n\n______________________________                         ______________________________\n[Course Faculty / Guide]                                               [HOD]\nDept. of ISE, GAT                                                        Professor and Head, Dept. of ISE, GAT")

    doc.add_page_break()
    doc.add_heading("ABSTRACT", level=1)
    _body(doc, f"A duplicate medical-question detector was implemented to help a health question-and-answer portal identify incoming questions that may already have been answered. The system used the Curai Medical Question Pairs dataset, containing {r.get('dataset_rows_clean', 'N/A')} cleaned labelled pairs; a reproducible subset of {r.get('subset_used', 'N/A')} pairs was selected with random seed {r.get('seed', 42)} and separated into training, validation and held-out test sets of {splits['train']}, {splits['validation']} and {splits['test']} examples. A custom GPT-2 Small architecture was implemented in PyTorch and initialized with pretrained GPT-2 weights. A binary classification head was added, and the final transformer block and final layer normalization were fine-tuned with AdamW and cross-entropy loss. On the held-out test set, accuracy changed from {_fmt_pct(before.get('accuracy'))} before fine-tuning to {_fmt_pct(after.get('accuracy'))} after fine-tuning, while F1 changed from {_fmt_pct(before.get('f1'))} to {_fmt_pct(after.get('f1'))}. The application accepts two new questions and returns SIMILAR or DIFFERENT with a model confidence score. The prototype is intended only for duplicate-question triage and not for diagnosis or patient-care decisions.")
    _body(doc, "Keywords: GPT-2, medical question pairs, binary classification, fine-tuning, PyTorch, Gradio.")

    doc.add_page_break()
    doc.add_heading("CHAPTER 1 — INTRODUCTION", level=1)
    doc.add_heading("1.1 Problem Statement", level=2)
    _body(doc, "A health Q&A portal receives new patient questions that may repeat questions already answered. Manual duplicate detection is time-consuming, while simple keyword overlap can mistake related but different medical questions for duplicates. This project predicts whether two medical questions express the same underlying intent so a portal can flag candidates for human review.")
    doc.add_heading("1.2 Objectives", level=2)
    _bullet(doc, "Build a binary question-pair classifier using the course-style GPT-2 Small architecture and pretrained weights.")
    _bullet(doc, "Prepare a reproducible train/validation/test split and report held-out accuracy, F1 and the confusion matrix before and after fine-tuning.")
    _bullet(doc, "Provide an interactive application for entering two previously unseen questions and displaying a prediction.")
    _bullet(doc, "Save and reload the trained checkpoint and generate plots and a report from measured results.")
    doc.add_heading("1.3 Category and Approach", level=2)
    _body(doc, "Category: MP-C classification fine-tuning. A two-class head is used because the output is a discrete similar/different label. The workflow uses GPT-2 tokenization and data loading, the course-style GPTModel architecture, pretrained weight loading, and the fine-tuning/evaluation concepts covered in the course.")

    doc.add_heading("CHAPTER 2 — DATASET AND PREPROCESSING", level=1)
    doc.add_heading("2.1 Dataset Description", level=2)
    _body(doc, "The Curai Medical Question Pairs dataset contains 3,048 doctor-labelled pairs based on patient questions. Label 1 denotes a similar-intent pair and label 0 denotes a different-intent pair. It is challenging by design: negative examples can share keywords while requiring a different answer, and positive examples can differ substantially in wording.")
    sample_records = r.get("dataset_samples", [])[:2]
    _table(doc, ["Sample question pair", "Label"], [[f"Q1: {x.get('question_1', '')}\nQ2: {x.get('question_2', '')}", "SIMILAR" if int(x.get("label", 0)) == 1 else "DIFFERENT"] for x in sample_records])
    _table(doc, ["Item", "Details"], [
        ["Source", "Curai MQP: https://github.com/curai/medical-question-pair-dataset"],
        ["Published size", "3,048 question pairs"],
        ["Cleaned rows", r.get("dataset_rows_clean", "N/A")],
        ["Subset used", r.get("subset_used", "N/A")],
        ["Label mapping", "0 = DIFFERENT; 1 = SIMILAR"],
        ["Train / validation / test", f"{splits['train']} / {splits['validation']} / {splits['test']}"],
        ["Tokenizer / maximum length", f"GPT-2 BPE (tiktoken) / {r.get('max_length', 96)} tokens"],
        ["Random seed", r.get("seed", 42)],
    ])
    doc.add_heading("2.2 Preprocessing Steps", level=2)
    _bullet(doc, "Downloaded the source CSV and parsed its headerless four-column format: doctor ID, question 1, question 2 and label.")
    _bullet(doc, "Trimmed whitespace, removed empty question pairs and removed exact duplicate question-pair/label rows; labels were validated as binary.")
    _bullet(doc, f"Selected a stratified subset of {r.get('subset_used', 'N/A')} examples, then created the recorded 70/15/15 train/validation/test split with seed {r.get('seed', 42)}.")
    _add_image(doc, OUT / "data_distribution.png", "Figure 2.1: Class distribution of the sampled dataset", width=4.4)
    doc.add_heading("2.3 Tokenization and Data Loader", level=2)
    _body(doc, f"Each input is formatted as Question 1 followed by Question 2 and encoded with the GPT-2 BPE tokenizer through tiktoken. The sequence is truncated to a maximum of {r.get('max_length', 96)} tokens while reserving token capacity for both questions. A custom PyTorch Dataset returns input token IDs and labels; the collate function pads each batch using GPT-2 token ID 50256 and records each sequence length. The classifier uses the logits at the final non-padding token.")

    doc.add_heading("CHAPTER 3 — MODEL AND TRAINING SETUP", level=1)
    doc.add_heading("3.1 Model Architecture", level=2)
    _body(doc, "The model uses a custom GPT-2 Small backbone with token and positional embeddings, 12 transformer blocks, multi-head causal self-attention, feed-forward layers, final layer normalization and a new two-class linear head. Pretrained GPT-2 weights are mapped into the PyTorch modules; no Hugging Face pipeline, Trainer or AutoModel class is used for the core model.")
    _table(doc, ["Parameter", "Value"], [
        ["Base model", "GPT-2 Small pretrained (124M)"],
        ["Vocabulary / context", "50,257 / 1,024 tokens"],
        ["Embedding dimension", "768"],
        ["Attention heads / blocks", "12 / 12"],
        ["Dropout", "0.0"],
        ["QKV bias", "Enabled"],
        ["Classification head", "Linear(768 → 2)"],
        ["Trainable layers", "Final transformer block + final layer norm + new head"],
        ["Total parameters", f"{int(r.get('total_parameters', 0)):,}"],
        ["Trainable parameters", f"{int(r.get('trainable_parameters', 0)):,}"],
    ])
    doc.add_heading("3.2 Changes Made for This Application", level=2)
    _body(doc, "The language-model output head was replaced by a new binary classification head. The last transformer block, final layer normalization and classifier were unfrozen, while the earlier pretrained blocks were frozen to reduce computation. The final real token representation summarises the pair context for the binary decision.")
    doc.add_heading("3.3 Training Settings", level=2)
    _table(doc, ["Setting", "Value"], [
        ["Optimizer", r.get("optimizer", "AdamW")],
        ["Learning rate / weight decay", f"{r.get('learning_rate', 'N/A')} / {r.get('weight_decay', 'N/A')}"],
        ["Batch size", r.get("batch_size", "N/A")],
        ["Epochs", r.get("epochs", "N/A")],
        ["Random seed", r.get("seed", 42)],
        ["Hardware", f"{r.get('device', 'unknown')}"],
        ["Training time", f"{float(r.get('training_seconds', 0)):.1f} seconds (measured for the training loop)"],
    ])

    doc.add_heading("CHAPTER 4 — RESULTS AND EVALUATION", level=1)
    doc.add_heading("4.1 Before and After Fine-Tuning", level=2)
    _table(doc, ["Metric", "Before", "After"], [
        ["Test accuracy", _fmt_pct(before.get("accuracy")), _fmt_pct(after.get("accuracy"))],
        ["Test F1", _fmt_pct(before.get("f1")), _fmt_pct(after.get("f1"))],
        ["Test precision", _fmt_pct(before.get("precision")), _fmt_pct(after.get("precision"))],
        ["Test recall", _fmt_pct(before.get("recall")), _fmt_pct(after.get("recall"))],
        ["Cross-entropy loss", _fmt_num(before.get("loss")), _fmt_num(after.get("loss"))],
    ])
    _body(doc, "The before values measure the pretrained backbone with a newly initialized classification head before optimization. The after values measure the same held-out test split after fine-tuning. The chart below shows the observed train and validation loss per epoch; interpretation should be based on the recorded curve rather than assumed improvement.")
    _add_image(doc, OUT / "loss_curve.png", "Figure 4.1: Training and validation loss curve", width=5.3)
    _add_image(doc, OUT / "confusion_matrix.png", "Figure 4.2: Confusion matrix on the held-out test split", width=4.2)
    doc.add_heading("4.2 Project-Specific Tasks", level=2)
    _table(doc, ["Task", "Implementation", "Recorded result"], [
        ["Design pair input and fine-tune a binary head", "Custom GPT-2 model, BPE tokenizer, frozen early blocks", f"{int(r.get('trainable_parameters', 0)):,} trainable parameters"],
        ["Report accuracy and F1", "Held-out test evaluation before and after training", f"Accuracy {_fmt_pct(before.get('accuracy'))} → {_fmt_pct(after.get('accuracy'))}; F1 {_fmt_pct(before.get('f1'))} → {_fmt_pct(after.get('f1'))}"],
        ["Test similar-looking but meaning-different pairs", "Held-out predictions plus five manually written examples", f"{sum(1 for x in r.get('demo_pairs', []) if x.get('correct'))}/5 demo examples matched expected labels"],
    ])
    doc.add_heading("4.3 Error Analysis", level=2)
    errs = r.get("error_examples", [])
    if errs:
        rows = []
        for e in errs[:3]:
            qpair = f"Q1: {e.get('question_1','')}\nQ2: {e.get('question_2','')}"
            rows.append([qpair, "SIMILAR" if int(e.get("expected", 0)) == 1 else "DIFFERENT", "SIMILAR" if int(e.get("predicted", 0)) == 1 else "DIFFERENT", "Possible shared keywords with different intent, medical details, negation or insufficient context."])
        _table(doc, ["Input pair", "Expected", "Prediction", "Possible reason"], rows)
    else:
        _body(doc, "No misclassified test examples were recorded in the selected sample. This does not establish perfect generalisation; additional examiner-written examples should be checked during the demo.")

    doc.add_heading("CHAPTER 5 — APPLICATION", level=1)
    doc.add_heading("5.1 Application Interface", level=2)
    _body(doc, "A Gradio web interface runs locally. Users enter two questions and press Compare questions. The app loads outputs/duplicate_detector.pt if available, uses the custom GPTModel and tokenizer, and displays SIMILAR or DIFFERENT with a confidence score. The application warns that results are not medical certainty and that a clinician should review duplicate candidates.")
    shot = OUT / "application_screenshot.png"
    _add_image(doc, shot, "Figure 5.1: Screenshot of the running application", width=5.8)
    if not shot.exists():
        _body(doc, "To complete Figure 5.1, run run_demo.bat, enter a real pair, take a screenshot, save it as outputs/application_screenshot.png, and run build_report.py again.")
    doc.add_heading("5.2 Sample Inputs and Outputs", level=2)
    demo_rows = []
    for i, d in enumerate(r.get("demo_pairs", [])[:5], start=1):
        qpair = f"Q1: {d.get('question_1','')}\nQ2: {d.get('question_2','')}"
        demo_rows.append([i, qpair, d.get("output", "N/A"), "Yes" if d.get("correct") else "No"])
    _table(doc, ["#", "Input", "Output", "Correct?"], demo_rows)

    doc.add_heading("CHAPTER 6 — CONCLUSION AND LIMITATIONS", level=1)
    _body(doc, f"The project implemented a local duplicate medical-question detector using a custom GPT-2 Small architecture and a binary classification head. Training used {r.get('subset_used', 'N/A')} pairs and a fixed random seed. On the held-out test split, accuracy was {_fmt_pct(before.get('accuracy'))} before fine-tuning and {_fmt_pct(after.get('accuracy'))} afterwards; F1 was {_fmt_pct(before.get('f1'))} and {_fmt_pct(after.get('f1'))}, respectively. The application loads the saved checkpoint and accepts unseen input pairs. Performance claims are limited to this particular subset and test split.")
    doc.add_heading("6.1 Limitations and Ethical Considerations", level=2)
    _bullet(doc, "A small sampled dataset may not represent all types of medical questions, writing styles or populations.")
    _bullet(doc, "The classifier can miss negation, dosage differences, age/sex context and other details that change medical meaning.")
    _bullet(doc, "This is not a clinical decision system; it should never triage, diagnose or recommend treatment without qualified human review.")
    doc.add_heading("6.2 Future Enhancements", level=2)
    _bullet(doc, "Compare additional frozen/unfrozen layer strategies and tune the threshold on a larger validation set.")
    _bullet(doc, "Evaluate with a domain-specific sentence-pair model and conduct expert review before any clinical use.")

    doc.add_heading("REFERENCES", level=1)
    for ref in [
        "[1] Curai, “Medical Question Pairs (MQP) Dataset,” GitHub. Available: https://github.com/curai/medical-question-pair-dataset.",
        "[2] A. Radford et al., “Language Models are Unsupervised Multitask Learners,” OpenAI Technical Report, 2019.",
        "[3] S. Raschka, Build a Large Language Model (From Scratch). Manning Publications, 2024. Source code: https://github.com/rasbt/LLMs-from-scratch.",
        "[4] OpenAI Community, “GPT-2 model weights,” Hugging Face. Available: https://huggingface.co/openai-community/gpt2.",
    ]:
        _body(doc, ref)
    doc.add_heading("AI Tool Declaration", level=2)
    _body(doc, "An AI coding assistant was used to help draft code scaffolding and documentation. The student must review, test and understand the final implementation and follow the course's academic-integrity policy.")

    doc.add_heading("APPENDIX A — CODE, NOTEBOOK LINK AND COMPLETE SOURCE", level=1)
    doc.add_heading("A.1 Links", level=2)
    _body(doc, "Complete code repository (GitHub): [INSERT PUBLIC GITHUB REPOSITORY URL AFTER UPLOADING]\nNotebook link (Colab): https://colab.research.google.com/github/YOUR_GITHUB_USERNAME/MP-C42-Duplicate-Medical-Question-Detector/blob/main/MP-C42_Duplicate_Medical_Question_Detector.ipynb\nMain notebook/file: MP-C42_Duplicate_Medical_Question_Detector.ipynb; main training script: train.py.\nPython/library versions: enter the versions printed by your final environment before submission.")
    doc.add_heading("A.2 Steps to Run", level=2)
    _bullet(doc, "Run setup.bat to create the virtual environment and install PyTorch plus project libraries.")
    _bullet(doc, "Run run_train.bat to download data and GPT-2 weights, train, evaluate, save the checkpoint, and create results and plots.")
    _bullet(doc, "Run run_demo.bat to launch the local Gradio application; use run_offline_demo.bat only for an emergency interface check.")
    _bullet(doc, "Do not upload the checkpoint/model weight files to GitHub; share the code, README, notebook, results and figures.")
    _add_source_appendix(doc)
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(REPORT)
    print(f"Report created: {REPORT}")
    return REPORT


if __name__ == "__main__":
    build_report(ROOT, require_results=True)
